#!/usr/bin/env python3
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).resolve().parent
D=json.loads((P/'R23_DECISION_REGISTER.json').read_text())
M=json.loads((P/'R23_B01_B18_MATRIX.json').read_text())
doc=Document(); sec=doc.sections[0]
sec.page_height,sec.page_width=Inches(11.7),Inches(8.3)
sec.left_margin=Inches(.78);sec.right_margin=Inches(.78);sec.top_margin=Inches(.79);sec.bottom_margin=Inches(.72)
sec.header_distance=Inches(.3);sec.footer_distance=Inches(.36)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Aptos';normal.font.size=Pt(9);normal.font.color.rgb=RGBColor(38,50,65)
normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.12
for name,size,color,after in [('Title',23,(21,43,66),10),('Heading 1',15,(21,43,66),7),('Heading 2',11,(29,70,95),5)]:
 st=styles[name];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(*color)
 st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(after)
if 'Small Note' not in styles:
 st=styles.add_style('Small Note',WD_STYLE_TYPE.PARAGRAPH);st.base_style=normal;st.font.size=Pt(8);st.font.color.rgb=RGBColor(88,101,115);st.paragraph_format.space_after=Pt(3)
header=sec.header.paragraphs[0]
header.style='Small Note';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
header.add_run('COMESTRE  /  MESTRE-SIDE  /  DMG R23').bold=True
footer=sec.footer.paragraphs[0];footer.style='Small Note';footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('MINUTA NÃO CANÔNICA  •  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def para(text='',style=None):return doc.add_paragraph(text,style=style)
def keyparagraph(label,value):
 p=doc.add_paragraph();p.paragraph_format.keep_together=True
 r=p.add_run(label+'  ');r.bold=True;r.font.color.rgb=RGBColor(28,72,93)
 p.add_run(value);return p

def table_headers(table,headers):
 for cell,value in zip(table.rows[0].cells,headers):
  cell.text=value
  tcPr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF3');tcPr.append(sh)
  for par in cell.paragraphs:
   for r in par.runs:r.bold=True;r.font.size=Pt(8)
 for row in table.rows:
  for cell in row.cells:
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(2)
    for run in par.runs:run.font.size=Pt(8)

p=doc.add_paragraph(style='Small Note');p.add_run('OP-DMG-CORPUS-EMP-01   •   DIR-MESTRE-DMG-R23-DECISION-20261009-001')
doc.add_heading('Dossiê de decisões do Mestre',0)
p=para('R23  |  B11 — Crítico  •  B12 — Veneno  •  B15 — Prova empírica');p.style='Subtitle'
keyparagraph('SITUAÇÃO','Cinco decisões preparadas. Zero respostas registradas. Não são regras aprovadas.')
keyparagraph('VALIDAÇÃO','R22 byte-verificado; IBX-76 consultada, Estado NOVA, sem ACK/instância writer.')
keyparagraph('BASELINE','B02 1/18 PASS documental   |   STRICT 0/2   |   I0003 0/9 originais   |   NAC-PASS HOLD')
para('Este documento é exclusivo do Mestre. As recomendações servem para comparar riscos e consequências; o campo de escolha deliberadamente fica em branco. A decisão expressa ainda exige NAC-NORM-01, verificação normativa, escritor legítimo, readback e gates aplicáveis antes de alterar o cânone.','Small Note')
doc.add_heading('Folha de apreciação',1)
t=doc.add_table(rows=1, cols=4);t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
table_headers(t,['DECISÃO','ESCOPO','PROPOSTA','ESCOLHA'])
for d in D['decisions']:
 row=t.add_row().cells
 row[0].text=d['id'].replace('DEC-','')
 row[1].text=d['title'].split(' — ')[-1]
 row[2].text=d['recommendation_option']+' (não aprovada)'
 row[3].text='__________'
for r in t.rows:
 for c in r.cells:
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(1)
   for run in p.runs:run.font.size=Pt(7.7)
para('Escolhas: A1–A3; B1–B3; C1–C3; D1–D3; E1–E3, ou texto livre. As alternativas abaixo estão detalhadas com custos, exemplos e riscos.','Small Note')
keyparagraph('ESCOPO / ÉPOCA AUTORIZADOS','_____________________________________________________')
keyparagraph('OBSERVAÇÕES DO MESTRE','_______________________________________________________')
doc.add_heading('Controle de continuidade',1)
keyparagraph('Preservado','R22 original e cadeia R01–R21; fórmulas e oráculo histórico R13 inalterados.')
keyparagraph('Pendente','Nenhum dos cinco itens é vigente por este documento; B11/B12 permanecem PARTIAL e B15 depende de evento real.')
keyparagraph('Não executar','Nenhuma escrita Party/O Primeiro, ficha, DMG-RANK, Notion ou VLAB a partir desta minuta.')

def page_decision(d):
 doc.add_page_break()
 p=doc.add_paragraph(style='Small Note');p.add_run('DECISÃO RESERVADA AO MESTRE  •  '+d['criterion']+'  •  '+d['id'])
 doc.add_heading(d['title'],0)
 keyparagraph('REGRA EXISTENTE',d['original'])
 keyparagraph('CONFLITO / OMISSÃO',d['gap'])
 doc.add_heading('Alternativas em apreciação',2)
 for o in d['choices']:
  p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
  r=p.add_run(o['option_id']+'  '+o['title']);r.bold=True;r.font.color.rgb=RGBColor(29,70,95)
  para(o['mechanic'])
  p=para('Consequência: '+o['risk'],style='Small Note');p.paragraph_format.space_after=Pt(5)
 doc.add_heading('Parecer técnico — não é aprovação',2)
 keyparagraph('RECOMENDAÇÃO',d['recommended'])
 keyparagraph('IMPACTO / EXEMPLO',d['worked'])
 keyparagraph('CONTRAPROVA',d['negative'])
 keyparagraph('FALTA PARA PASS',d['remaining'])
 p=para('FONTES: '+', '.join(D['source_count'] and [b for b in d['basis']] or []),'Small Note')
 for par in [p]:par.paragraph_format.space_after=Pt(9)
 para('ESCOLHA EXPRESSA DO MESTRE: ________________     JUSTIFICATIVA: ____________________')

for d in D['decisions']:page_decision(d)

doc.add_page_break();doc.add_heading('B15 — Prova empírica mínima',0)
keyparagraph('CONTRATO ORIGINAL','ELEMENTAL_WEAKNESS; status PARTIAL. Regra de dano misto já é documentalmente suficiente.')
keyparagraph('EXEMPLO CANÔNICO','1d8 cortante +2d6 Fogo → 1d8 cortante +5d6 Fogo com Fraqueza a Fogo fonte-confirmada. Médias: 11,5 →22; ganho médio +10,5. O exemplo não é uma arma equipada.')
doc.add_heading('Evento prospectivo a recolher',2)
short=[
 'Mestre autoriza cena e identifica CM-000006 quando for execução oficial; registrar scene_id, epoch, effective_at e fonte congelada.',
 'Atacante tem ação EXISTENTE, legal, com componente físico e Fogo explicitamente separados. Alvo tem Fraqueza Fogo fonte-comprovada.',
 'Registrar alcance/LoS, ação, PA, Defesa/Reação, hit ou resolução específica de magia; fontes de RD, resistências e imunidades.',
 'Registrar cada dado físico, Fogo-base e Fogo-extra, ordem de mitigação por tag; HP/Alma/PA antes/depois; event_id/commit e recibo antiduplo.',
 'Refazer cálculo por leitura independente. Controle negativo: sem fraqueza ou sem tag Fogo, não aplica três dados extras.',
]
for x in short:doc.add_paragraph(x,style='List Bullet')
keyparagraph('ACEITE','PASS de B15 somente após evento REAL, fonte vigente, cálculo íntegro, readback independente e NAC-PASS. Não é combate STRICT inteiro.')
keyparagraph('NÃO EXECUTADO','Nenhum evento real recolhido nesta rodada. 0/2 STRICT permanece.')
para('O contrato integral de 37 campos e as etapas de aquisição constam do plano anexo no ZIP.','Small Note')

doc.add_page_break();doc.add_heading('Matriz científica e próximos atos',0)
t=doc.add_table(rows=1,cols=3);t.style='Table Grid'
table_headers(t,['CRITÉRIO','ESTADO R23','PRÓXIMO GATE MÍNIMO'])
for r in M['rows']:
 row=t.add_row().cells;row[0].text=r['criterion'];row[1].text=r['estado_R23'];
 s=r['R23_acceptance_blocker'];row[2].text=s[:135]+('…' if len(s)>135 else '')
for row in t.rows:
 for c in row.cells:
  for par in c.paragraphs:
   par.paragraph_format.space_after=Pt(1)
   for run in par.runs:run.font.size=Pt(7)

doc.add_heading('Fronteira de retomada',2)
para('1. Receber cinco decisões expressas. 2. Conferir precedência NAC-NORM-01/escopo/epoch antes de qualquer atualização. 3. Usar writer legítimo e readback + NAC-PASS. 4. Obter evento real B15, e eventos críticos e de veneno após aprovação. 5. Liu e I0003 somente mediante nova pista material.')
para('Nenhuma das cinco recomendações cria, por si só, um PASS adicional. Baseline final: 1/18; 0/2 STRICT; 0/9 originais.','Small Note')

# Prevent implicit local approval metadata in document.
doc.core_properties.title='DMG R23 — Dossiê decisório Mestre-side NÃO CANÔNICO'
doc.core_properties.subject='Decisões propostas B11 A/B/C e B12 A/B; coleta empírica B15'
doc.core_properties.author='Comestre / CM-000001 (pesquisa Mestre-side)'
doc.core_properties.keywords='MESTRE SIDE; DRAFT; R23'
doc.save(P/'DOSSIER_DECISORIO_R23.docx')
print('DOCX_BUILT',P/'DOSSIER_DECISORIO_R23.docx')
